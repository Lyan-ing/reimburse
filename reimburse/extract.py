from __future__ import annotations

import email
import html
import io
import os
import re
import shutil
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Callable
from xml.etree import ElementTree as ET

from .models import Material
from .ocr import OcrEngine
from .scanner import sha1_file

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp", ".gif"}
PDF_SUFFIXES = {".pdf"}
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".log"}
OFFICE_SUFFIXES = {".docx", ".xlsx", ".pptx"}
LEGACY_OFFICE_SUFFIXES = {".doc", ".xls", ".ppt"}


def _decode_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "utf-16", "big5"):
        try:
            text = data.decode(encoding)
            if text.count("\ufffd") <= max(2, len(text) // 100):
                return text
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _strip_html(text: str) -> str:
    text = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    text = html.unescape(text)
    return re.sub(r"[ \t]+", " ", text)


def _xml_text(data: bytes) -> str:
    try:
        root = ET.fromstring(data)
    except ET.ParseError:
        return _decode_bytes(data)
    texts: list[str] = []
    for elem in root.iter():
        tag = elem.tag.split("}")[-1].lower()
        if elem.text and elem.text.strip() and tag in {
            "textcode",
            "text",
            "title",
            "content",
            "name",
            "value",
            "remark",
            "memo",
        }:
            texts.append(elem.text.strip())
    if not texts:
        for elem in root.iter():
            if elem.text and elem.text.strip():
                texts.append(elem.text.strip())
    return "\n".join(texts)


def _ofd_text(path: Path) -> str:
    parts: list[str] = []
    try:
        with zipfile.ZipFile(path) as zf:
            names = [n for n in zf.namelist() if n.lower().endswith(".xml")]
            names.sort(key=lambda n: (0 if "content" in n.lower() else 1, n))
            for name in names:
                try:
                    text = _xml_text(zf.read(name))
                except Exception:
                    continue
                if text.strip():
                    parts.append(text)
    except (zipfile.BadZipFile, OSError) as exc:
        raise ValueError(f"OFD 文件无法读取: {exc}") from exc
    return "\n".join(parts)


def _docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if cell.text.strip():
                    cells.append(cell.text.strip())
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            parts.append(ws.title)
            for row in ws.iter_rows(max_row=200, values_only=True):
                values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if values:
                    parts.append(" | ".join(values))
            if len(parts) > 3000:
                break
    finally:
        wb.close()
    return "\n".join(parts)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
    return "\n".join(parts)


def _eml_text(path: Path) -> str:
    raw = path.read_bytes()
    msg = email.message_from_bytes(raw)
    parts = [str(msg.get("Subject", "")), str(msg.get("From", "")), str(msg.get("Date", ""))]
    if msg.is_multipart():
        for part in msg.walk():
            ctype = part.get_content_type()
            if ctype in {"text/plain", "text/html"}:
                payload = part.get_payload(decode=True) or b""
                text = _decode_bytes(payload)
                parts.append(_strip_html(text) if ctype == "text/html" else text)
    else:
        payload = msg.get_payload(decode=True) or b""
        text = _decode_bytes(payload)
        parts.append(_strip_html(text) if msg.get_content_type() == "text/html" else text)
    return "\n".join(p for p in parts if p.strip())


class Extractor:
    def __init__(self, use_ocr: bool = True) -> None:
        self.use_ocr = use_ocr
        self.ocr = OcrEngine()
        self.temp_dir = Path(tempfile.mkdtemp(prefix="reimburse_preview_"))

    def cleanup(self) -> None:
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def extract(self, path: Path) -> Material:
        material = Material(path=path)
        try:
            stat = path.stat()
            material.size_bytes = stat.st_size
            material.mtime = stat.st_mtime
        except OSError:
            material.error = "文件无法访问"
            return material
        material.source_hash = sha1_file(path)
        suffix = path.suffix.lower()
        try:
            if suffix in PDF_SUFFIXES:
                material.page_count, material.text, material.method = self._pdf(path)
            elif suffix == ".ofd":
                material.text = _ofd_text(path)
                material.method = "OFD文字"
                material.page_count = 1
            elif suffix in {".xml", ".xbrl"}:
                material.text = _xml_text(path.read_bytes())
                material.method = "XML文字"
                material.page_count = 1
            elif suffix == ".docx":
                material.text = _docx_text(path)
                material.method = "Word文字"
                material.page_count = 1
            elif suffix == ".xlsx":
                material.text = _xlsx_text(path)
                material.method = "Excel文字"
                material.page_count = 1
            elif suffix == ".pptx":
                material.text = _pptx_text(path)
                material.method = "PPT文字"
                material.page_count = 1
            elif suffix == ".eml":
                material.text = _eml_text(path)
                material.method = "邮件文字"
                material.page_count = 1
            elif suffix in {".html", ".htm"}:
                material.text = _strip_html(_decode_bytes(path.read_bytes()))
                material.method = "网页文字"
                material.page_count = 1
            elif suffix in TEXT_SUFFIXES:
                material.text = _decode_bytes(path.read_bytes())
                material.method = "文本"
                material.page_count = 1
            elif suffix in IMAGE_SUFFIXES:
                material.page_count = 1
                material.preview_pages = [str(path)]
                if self.use_ocr:
                    text, error = self.ocr.recognize(path)
                    material.text = text
                    material.method = "图片OCR"
                    material.error = error
                else:
                    material.method = "图片（未OCR）"
            elif suffix in LEGACY_OFFICE_SUFFIXES:
                material.method = "旧版Office文件"
                material.error = "旧版 Office 格式无法读取文字，仅登记材料"
                material.page_count = 1
            else:
                material.method = "未解析"
                material.page_count = 1
        except Exception as exc:
            material.error = str(exc)[:300]
            material.method = "提取失败"
        material.text = (material.text or "").strip()
        return material

    def _pdf(self, path: Path) -> tuple[int, str, str]:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        page_count = len(reader.pages)
        texts: list[str] = []
        for page in reader.pages[:8]:
            try:
                texts.append(page.extract_text() or "")
            except Exception:
                continue
        text = "\n".join(texts).strip()
        method = "PDF文字"
        if self._needs_pdf_ocr(text):
            if self.use_ocr:
                rendered = self.render_pdf_page(path, 0, scale=2.4)
                if rendered:
                    ocr_text, error = self.ocr.recognize(rendered)
                    if ocr_text:
                        text = ocr_text
                        method = "PDF扫描OCR"
                    if error and not ocr_text:
                        method = "PDF扫描（OCR失败）"
            else:
                method = "PDF扫描（未OCR）"
        return page_count, text, method

    @staticmethod
    def _needs_pdf_ocr(text: str) -> bool:
        cleaned = re.sub(r"\s+", "", text)
        if len(cleaned) < 12:
            return True
        keywords = ("发票", "价税合计", "行程单", "会议", "通知", "账单", "交易", "铁路", "航空")
        if any(word in cleaned for word in keywords):
            return False
        # English notices/emails are fine without OCR if there is real text.
        if len(cleaned) >= 120 and re.search(r"[A-Za-z]{4,}", cleaned):
            return False
        return len(cleaned) < 80

    def render_pdf_page(self, path: Path, page_index: int = 0, scale: float = 2.4) -> Path | None:
        try:
            import pypdfium2 as pdfium
        except Exception:
            return None
        try:
            pdf = pdfium.PdfDocument(str(path))
            try:
                if page_index >= len(pdf):
                    return None
                page = pdf[page_index]
                bitmap = page.render(scale=scale)
                pil_image = bitmap.to_pil()
                out = self.temp_dir / f"{sha1_file(path)[:12]}_{page_index}.png"
                if not out.exists():
                    pil_image.save(out, "PNG")
                return out
            finally:
                pdf.close()
        except Exception:
            return None

    def render_pdf_pages(self, path: Path, max_pages: int = 8, scale: float = 2.0) -> list[Path]:
        try:
            import pypdfium2 as pdfium
        except Exception:
            return []
        outputs: list[Path] = []
        try:
            pdf = pdfium.PdfDocument(str(path))
            try:
                for index in range(min(len(pdf), max_pages)):
                    page = pdf[index]
                    bitmap = page.render(scale=scale)
                    pil_image = bitmap.to_pil()
                    out = self.temp_dir / f"{sha1_file(path)[:12]}_p{index}.png"
                    if not out.exists():
                        pil_image.save(out, "PNG")
                    outputs.append(out)
            finally:
                pdf.close()
        except Exception:
            return outputs
        return outputs


def preview_paths_for(material: Material) -> list[Path]:
    if material.preview_pages:
        paths = [Path(p) for p in material.preview_pages]
        if all(p.exists() for p in paths):
            return paths
    suffix = material.suffix
    if suffix in IMAGE_SUFFIXES:
        return [material.path]
    return []
