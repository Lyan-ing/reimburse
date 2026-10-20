from __future__ import annotations

import os
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont, ImageOps

from .models import (
    SECTION_ORDER,
    ExpenseGroup,
    Material,
    ScanResult,
)

PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
MARGIN_CM = 1.25
USABLE_WIDTH_CM = PAGE_WIDTH_CM - 2 * MARGIN_CM
USABLE_HEIGHT_CM = PAGE_HEIGHT_CM - 2 * MARGIN_CM
MAX_BLOCK_HEIGHT_CM = USABLE_HEIGHT_CM - 1.45
COLUMN_WIDTH_CM = USABLE_WIDTH_CM / 2
FULL_IMAGE_WIDTH_CM = USABLE_WIDTH_CM - 0.15
HALF_IMAGE_WIDTH_CM = COLUMN_WIDTH_CM - 0.15
IMAGE_GAP_CM = 0.10


@dataclass
class ImageBlock:
    images: list[tuple[Path, float, float]] = field(default_factory=list)
    width_cm: float = FULL_IMAGE_WIDTH_CM
    height_cm: float = 0.0
    full_width: bool = True
    label: str = ""


def _sanitize_filename(name: str, limit: int = 48) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", name)
    name = name.strip(" ._-")
    name = re.sub(r"\s+", " ", name)
    if len(name) > limit:
        name = name[:limit].rstrip(" ._-")
    return name or "报销事项"


def _date_part(start: date | None, end: date | None) -> str:
    if not start and not end:
        return ""
    if start and end:
        if start.year == end.year:
            return f"{start:%Y%m%d}-{end:%m%d}"
        return f"{start:%Y%m%d}-{end:%Y%m%d}"
    value = start or end
    return f"{value:%Y%m%d}" if value else ""


def default_filename(result: ScanResult) -> str:
    event = _sanitize_filename(result.event_name or result.folder.name)
    dates = _date_part(result.date_start, result.date_end)
    total = result.total
    amount = f"{total:,.2f}元"
    parts = ["报销材料", event]
    if dates:
        parts.append(dates)
    parts.append(f"合计{amount}")
    return "_".join(p for p in parts if p) + ".docx"


def unique_output_path(path: Path) -> Path:
    if not path.exists():
        return path
    for index in range(2, 1000):
        candidate = path.with_name(f"{path.stem}({index}){path.suffix}")
        if not candidate.exists():
            return candidate
    return path


def _font(size: int) -> ImageFont.FreeTypeFont:
    for name in ("msyh.ttc", "simhei.ttf", "simsun.ttc", "arial.ttf"):
        path = Path(r"C:\Windows\Fonts") / name
        if path.exists():
            try:
                return ImageFont.truetype(str(path), size)
            except Exception:
                continue
    return ImageFont.load_default()


def _wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    lines: list[str] = []
    for raw in text.splitlines() or [""]:
        current = ""
        for char in raw:
            trial = current + char
            if draw.textlength(trial, font=font) <= max_width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = char
        if current:
            lines.append(current)
    return lines


def make_placeholder(material: Material, output: Path, reason: str = "") -> Path:
    width, height = 1654, 2339
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_title = _font(46)
    font_body = _font(30)
    draw.rectangle((60, 60, width - 60, height - 60), outline="#C9D3E0", width=4)
    y = 260
    title = "材料预览不可用"
    draw.text((120, y), title, font=font_title, fill="#1F3864")
    y += 100
    draw.text((120, y), "请将原件一并提交，或转换为 PDF/图片后重新生成。", font=font_body, fill="#555555")
    y += 90
    lines = [
        f"文件名：{material.name}",
        f"材料类型：{material.section_name} / {material.role_name}",
        f"格式：{material.suffix or '未知'}",
        f"识别方式：{material.method}",
    ]
    if material.amount is not None:
        lines.append(f"票面金额：¥{material.amount:,.2f}")
    if material.business_date:
        lines.append(f"业务日期：{material.business_date.isoformat()}")
    if reason:
        lines.append(f"原因：{reason}")
    elif material.error:
        lines.append(f"原因：{material.error}")
    for line in lines:
        for wrapped in _wrap_text(draw, line, font_body, width - 240):
            draw.text((120, y), wrapped, font=font_body, fill="#333333")
            y += 48
    y += 40
    draw.line((120, y, width - 120, y), fill="#D9E2F3", width=2)
    y += 40
    for line in _wrap_text(draw, "提示：OFD、旧版 Office 等格式暂不支持自动排版预览，但识别到的金额仍会计入总额。", font_body, width - 240):
        draw.text((120, y), line, font=font_body, fill="#7A7A7A")
        y += 46
    image.save(output, "PNG")
    return output


class PreviewRenderer:
    def __init__(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp(prefix="reimburse_doc_"))
        self._cache: dict[tuple[str, int], Path] = {}

    def cleanup(self) -> None:
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def _normalized_image(self, material: Material) -> Path | None:
        try:
            key = (str(material.path), 0)
            if key in self._cache:
                return self._cache[key]
            with Image.open(material.path) as img:
                img = ImageOps.exif_transpose(img)
                if img.mode in {"RGBA", "LA", "P"}:
                    background = Image.new("RGB", img.size, "white")
                    if img.mode == "P":
                        img = img.convert("RGBA")
                    background.paste(img, mask=img.getchannel("A") if "A" in img.getbands() else None)
                    img = background
                elif img.mode != "RGB":
                    img = img.convert("RGB")
                max_width = 1800
                if img.width > max_width:
                    ratio = max_width / img.width
                    img = img.resize((max_width, int(img.height * ratio)), Image.LANCZOS)
                out = self.temp_dir / f"{material.source_hash[:12] or 'img'}_{abs(hash(material.uid)) & 0xFFFFFF}.png"
                img.save(out, "PNG")
                self._cache[key] = out
                return out
        except Exception:
            return None

    def pages(self, material: Material) -> list[Path]:
        if material.suffix in {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp", ".gif"}:
            normalized = self._normalized_image(material)
            return [normalized] if normalized else [make_placeholder(material, self.temp_dir / f"ph_{abs(hash(material.uid))}.png", "图片无法读取")]
        if material.suffix == ".pdf":
            max_pages = 6 if material.role == "notice" else 1
            return self._render_pdf(material, max_pages)
        return [make_placeholder(material, self.temp_dir / f"ph_{abs(hash(material.uid))}.png")]

    def _render_pdf(self, material: Material, max_pages: int) -> list[Path]:
        try:
            import pypdfium2 as pdfium

            key = (str(material.path), max_pages)
            if key in self._cache:
                return [self._cache[key]]
            pdf = pdfium.PdfDocument(str(material.path))
            outputs: list[Path] = []
            try:
                for index in range(min(len(pdf), max_pages)):
                    page = pdf[index]
                    bitmap = page.render(scale=2.2)
                    pil_image = bitmap.to_pil()
                    if pil_image.width > 1900:
                        ratio = 1900 / pil_image.width
                        pil_image = pil_image.resize((1900, int(pil_image.height * ratio)), Image.LANCZOS)
                    out = self.temp_dir / f"{material.source_hash[:12] or 'pdf'}_p{index}.png"
                    pil_image.save(out, "PNG")
                    outputs.append(out)
            finally:
                pdf.close()
            if outputs:
                self._cache[key] = outputs[0]
            return outputs
        except Exception as exc:
            return [make_placeholder(material, self.temp_dir / f"ph_{abs(hash(material.uid))}.png", f"PDF 渲染失败：{exc}")]


def _image_size(path: Path) -> tuple[float, float]:
    try:
        with Image.open(path) as img:
            return float(img.width), float(img.height)
    except Exception:
        return 4.0, 3.0


def _make_stack_block(images: list[Path], width_cm: float, label: str = "") -> ImageBlock:
    entries: list[tuple[Path, float, float]] = []
    total_height = 0.0
    for path in images:
        width_px, height_px = _image_size(path)
        display_width = width_cm
        display_height = display_width * height_px / max(width_px, 1)
        entries.append((path, display_width, display_height))
        total_height += display_height
    if len(entries) > 1:
        total_height += IMAGE_GAP_CM * (len(entries) - 1)
    if total_height > MAX_BLOCK_HEIGHT_CM:
        scale = MAX_BLOCK_HEIGHT_CM / total_height
        entries = [(p, w * scale, h * scale) for p, w, h in entries]
        total_height = MAX_BLOCK_HEIGHT_CM
    return ImageBlock(images=entries, width_cm=max((e[1] for e in entries), default=width_cm), height_cm=total_height, full_width=True, label=label)


def _make_half_block(path: Path, label: str = "") -> ImageBlock:
    width_px, height_px = _image_size(path)
    display_width = HALF_IMAGE_WIDTH_CM
    display_height = display_width * height_px / max(width_px, 1)
    if display_height > MAX_BLOCK_HEIGHT_CM:
        scale = MAX_BLOCK_HEIGHT_CM / display_height
        display_width *= scale
        display_height = MAX_BLOCK_HEIGHT_CM
    return ImageBlock(images=[(path, display_width, display_height)], width_cm=display_width, height_cm=display_height, full_width=False, label=label)


def _group_materials(group: ExpenseGroup) -> list[Material]:
    order = {"invoice": 0, "itinerary": 1, "toll": 2, "photo": 3, "notice": 4, "unknown": 5}
    materials = [m for m in group.materials if m.include and not m.duplicate]
    materials.sort(key=lambda m: (order.get(m.role, 9), m.business_date or date.max, m.name))
    return materials


def _blocks_for_group(group: ExpenseGroup, renderer: PreviewRenderer) -> list[ImageBlock]:
    materials = _group_materials(group)
    if not materials:
        return []
    if group.section == "taxi":
        paths: list[Path] = []
        for material in materials:
            if material.role == "payment":
                continue
            paths.extend(renderer.pages(material))
        if not paths:
            return []
        return [_make_stack_block(paths, FULL_IMAGE_WIDTH_CM, label=group.gid)]
    if group.section == "notice":
        blocks: list[ImageBlock] = []
        for material in materials:
            for page in renderer.pages(material):
                blocks.append(_make_stack_block([page], FULL_IMAGE_WIDTH_CM, label=group.gid))
        return blocks
    if group.section == "other" and group.subtype == "purchase":
        blocks: list[ImageBlock] = []
        if group.invoice:
            blocks.extend(_make_stack_block([p], FULL_IMAGE_WIDTH_CM, label=group.gid) for p in renderer.pages(group.invoice))
        for photo in group.photos:
            for page in renderer.pages(photo):
                blocks.append(_make_stack_block([page], FULL_IMAGE_WIDTH_CM, label=group.gid))
        if blocks:
            return blocks
    blocks: list[ImageBlock] = []
    for material in materials:
        if material.role == "payment":
            continue
        for page in renderer.pages(material):
            blocks.append(_make_stack_block([page], FULL_IMAGE_WIDTH_CM, label=group.gid))
    return blocks


def _pack_full_blocks(blocks: list[ImageBlock]) -> list[list[ImageBlock]]:
    pages: list[list[ImageBlock]] = []
    current: list[ImageBlock] = []
    used = 0.0
    for block in blocks:
        gap = IMAGE_GAP_CM if current else 0.0
        if current and used + block.height_cm + gap > USABLE_HEIGHT_CM:
            pages.append(current)
            current = []
            used = 0.0
            gap = 0.0
        current.append(block)
        used += block.height_cm + gap
    if current:
        pages.append(current)
    return pages


def _pack_half_blocks(blocks: list[ImageBlock]) -> list[list[list[ImageBlock]]]:
    pages: list[list[list[ImageBlock]]] = []
    rows: list[list[ImageBlock]] = []
    used = 0.0
    for block in blocks:
        if rows and len(rows[-1]) == 1:
            old_height = rows[-1][0].height_cm
            new_height = max(old_height, block.height_cm)
            proposed = used - old_height + new_height
            if proposed <= MAX_BLOCK_HEIGHT_CM:
                rows[-1].append(block)
                used = proposed
                continue
            pages.append(rows)
            rows = []
            used = 0.0
        gap = IMAGE_GAP_CM if rows else 0.0
        if rows and used + gap + block.height_cm > MAX_BLOCK_HEIGHT_CM:
            pages.append(rows)
            rows = []
            used = 0.0
            gap = 0.0
        rows.append([block])
        used += gap + block.height_cm
    if rows:
        pages.append(rows)
    return pages


def _set_cell_margins(cell, top=0, start=0, bottom=0, end=0) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borderless(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    existing = tbl_pr.find(qn("w:tblBorders"))
    if existing is not None:
        tbl_pr.remove(existing)
    tbl_pr.append(borders)


def _set_table_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(int(USABLE_WIDTH_CM / 2.54 * 1440)))
    width.set(qn("w:type"), "dxa")
    existing = tbl_pr.find(qn("w:tblW"))
    if existing is not None:
        tbl_pr.remove(existing)
    tbl_pr.append(width)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    cell_mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")
        cell_mar.append(node)
    tbl_pr.append(cell_mar)


def _no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def _add_picture_paragraph(cell, path: Path, width_cm: float, first: bool = False):
    if first:
        paragraph = cell.paragraphs[0]
    else:
        paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    return paragraph


def _add_full_block(table, block: ImageBlock) -> None:
    row = table.add_row()
    _no_split(row)
    cell = row.cells[0].merge(row.cells[1])
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    _set_cell_margins(cell)
    first = True
    for index, (path, width, height) in enumerate(block.images):
        if index > 0:
            spacer = cell.add_paragraph()
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.space_after = Pt(0)
            spacer.paragraph_format.line_spacing = 1
            spacer_run = spacer.add_run()
            spacer_run.font.size = Pt(1)
        _add_picture_paragraph(cell, path, width, first=first)
        first = False


def _add_half_row(table, blocks: list[ImageBlock]) -> None:
    row = table.add_row()
    _no_split(row)
    for index in range(2):
        cell = row.cells[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        _set_cell_margins(cell)
        if index < len(blocks):
            block = blocks[index]
            if block.images:
                path, width, _ = block.images[0]
                _add_picture_paragraph(cell, path, width, first=True)


def _new_page(doc: Document):
    section = doc.sections[0]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(COLUMN_WIDTH_CM)
    table.columns[1].width = Cm(COLUMN_WIDTH_CM)
    _set_table_layout(table)
    _set_table_borderless(table)
    return table


def _add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)


def build_document(result: ScanResult, output_path: Path) -> Path:
    output_path = unique_output_path(output_path)
    renderer = PreviewRenderer()
    try:
        doc = Document()
        section = doc.sections[0]
        section.page_width = Cm(PAGE_WIDTH_CM)
        section.page_height = Cm(PAGE_HEIGHT_CM)
        section.orientation = WD_ORIENT.PORTRAIT
        section.top_margin = Cm(MARGIN_CM)
        section.bottom_margin = Cm(MARGIN_CM)
        section.left_margin = Cm(MARGIN_CM)
        section.right_margin = Cm(MARGIN_CM)
        section.header_distance = Cm(0.5)
        section.footer_distance = Cm(0.5)

        normal = doc.styles["Normal"]
        normal.font.name = "宋体"
        normal.font.size = Pt(9)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = 1
        doc.core_properties.title = "报销材料"
        doc.core_properties.subject = result.event_name

        pages: list[tuple[str, list[ImageBlock], bool]] = []
        # Build full-width invoice pages by section, then payments, then notices.
        for section_name in ("traffic", "hotel", "taxi", "other"):
            section_groups = [g for g in result.groups if g.section == section_name]
            blocks: list[ImageBlock] = []
            for group in section_groups:
                blocks.extend(_blocks_for_group(group, renderer))
            if blocks:
                for page in _pack_full_blocks(blocks):
                    pages.append((section_name, page, True))

        # Payment screenshots use a two-column layout, ordered by their invoice group.
        payment_blocks: list[ImageBlock] = []
        for payment in result.ordered_payments:
            if not payment.include or payment.duplicate:
                continue
            for page_image in renderer.pages(payment):
                payment_blocks.append(_make_half_block(page_image, label=payment.group_id))
        if payment_blocks:
            for page in _pack_half_blocks(payment_blocks):
                flat = [block for row in page for block in row]
                pages.append(("payment", flat, False))

        notice_groups = [g for g in result.groups if g.section == "notice"]
        notice_blocks: list[ImageBlock] = []
        for group in notice_groups:
            notice_blocks.extend(_blocks_for_group(group, renderer))
        if notice_blocks:
            for page in _pack_full_blocks(notice_blocks):
                pages.append(("notice", page, True))

        if not pages:
            doc.add_paragraph("没有可输出的材料。")

        for index, (_, page_blocks, full_width) in enumerate(pages):
            if index > 0:
                _add_page_break(doc)
            table = _new_page(doc)
            if full_width:
                for block in page_blocks:
                    _add_full_block(table, block)
            else:
                # Repack flat payment blocks into rows of two.
                for start in range(0, len(page_blocks), 2):
                    _add_half_row(table, page_blocks[start : start + 2])

        doc.save(str(output_path))
        return output_path
    finally:
        renderer.cleanup()
